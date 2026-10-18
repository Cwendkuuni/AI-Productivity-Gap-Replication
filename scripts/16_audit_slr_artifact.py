#!/usr/bin/env python3

from pathlib import Path
from collections import Counter

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

XLSX = (
    ROOT
    / "data"
    / "raw"
    / "slr_artifact"
    / "SLR_all_included_studies.xlsx"
)

README = (
    ROOT
    / "data"
    / "raw"
    / "slr_artifact"
    / "Readme.md"
)

OUT_TABLE_DIR = ROOT / "results" / "tables"
OUT_LOG_DIR = ROOT / "results" / "logs"

OUT_TABLE_DIR.mkdir(parents=True, exist_ok=True)
OUT_LOG_DIR.mkdir(parents=True, exist_ok=True)

LOG_FILE = (
    OUT_LOG_DIR
    / "slr_artifact_audit.txt"
)

STRUCTURE_FILE = (
    OUT_TABLE_DIR
    / "slr_workbook_structure.csv"
)

COMPLETENESS_FILE = (
    OUT_TABLE_DIR
    / "slr_column_completeness.csv"
)

VALUES_FILE = (
    OUT_TABLE_DIR
    / "slr_value_inventory.csv"
)


def clean_text(value):
    if pd.isna(value):
        return ""

    return str(value).strip()


def shorten(value, max_len=150):
    value = clean_text(value)

    if len(value) <= max_len:
        return value

    return value[: max_len - 3] + "..."


def main():

    log_lines = []

    def log(message=""):
        print(message)
        log_lines.append(str(message))

    # =========================================================
    # Basic file checks
    # =========================================================

    log("=" * 80)
    log("SLR ARTIFACT AUDIT")
    log("=" * 80)
    log()

    log(f"Workbook: {XLSX}")
    log(f"Exists:   {XLSX.exists()}")

    if not XLSX.exists():
        raise FileNotFoundError(XLSX)

    log(
        f"Size:     "
        f"{XLSX.stat().st_size:,} bytes"
    )

    log()

    # =========================================================
    # Workbook structure via openpyxl
    # =========================================================

    wb = load_workbook(
        XLSX,
        read_only=True,
        data_only=False
    )

    log("WORKBOOK SHEETS")
    log("-" * 80)

    for sheet in wb.sheetnames:

        ws = wb[sheet]

        log(
            f"{sheet}: "
            f"{ws.max_row} rows x "
            f"{ws.max_column} columns"
        )

    log()

    # =========================================================
    # Read all sheets with pandas
    # =========================================================

    sheets = pd.read_excel(
        XLSX,
        sheet_name=None
    )

    structure_rows = []
    completeness_rows = []
    value_rows = []

    for sheet_name, df in sheets.items():

        log("=" * 80)
        log(f"SHEET: {sheet_name}")
        log("=" * 80)

        log(
            f"Shape: {df.shape[0]} rows x "
            f"{df.shape[1]} columns"
        )

        log()

        log("COLUMNS")
        log("-" * 80)

        for idx, col in enumerate(
            df.columns,
            start=1
        ):
            log(
                f"{idx:02d}. {col}"
            )

            structure_rows.append(
                {
                    "sheet": sheet_name,
                    "column_index": idx,
                    "column_name": str(col),
                    "rows": len(df),
                }
            )

        log()

        # -----------------------------------------------------
        # Completeness
        # -----------------------------------------------------

        log("COLUMN COMPLETENESS")
        log("-" * 80)

        for col in df.columns:

            non_missing = int(
                df[col].notna().sum()
            )

            missing = int(
                df[col].isna().sum()
            )

            pct_complete = (
                100.0 * non_missing / len(df)
                if len(df) > 0
                else 0.0
            )

            unique = int(
                df[col]
                .dropna()
                .astype(str)
                .nunique()
            )

            log(
                f"{col}: "
                f"{non_missing}/{len(df)} "
                f"non-missing "
                f"({pct_complete:.1f}%), "
                f"{unique} unique"
            )

            completeness_rows.append(
                {
                    "sheet": sheet_name,
                    "column": str(col),
                    "rows": len(df),
                    "non_missing": non_missing,
                    "missing": missing,
                    "complete_pct": round(
                        pct_complete,
                        2
                    ),
                    "unique_values": unique,
                }
            )

        log()

        # -----------------------------------------------------
        # First rows
        # -----------------------------------------------------

        log("FIRST 5 ROWS")
        log("-" * 80)

        if len(df) == 0:
            log("(empty sheet)")
        else:
            preview = (
                df.head(5)
                .fillna("")
                .astype(str)
            )

            log(
                preview.to_string(
                    index=False,
                    max_colwidth=40
                )
            )

        log()

        # -----------------------------------------------------
        # Value inventory
        # -----------------------------------------------------

        log("VALUE INVENTORY")
        log("-" * 80)

        for col in df.columns:

            series = (
                df[col]
                .dropna()
                .astype(str)
                .str.strip()
            )

            series = series[
                series != ""
            ]

            counts = Counter(series)

            # Print frequent values only when useful.
            if len(counts) <= 20:

                log()
                log(f"[{col}]")

                for value, count in (
                    counts.most_common(20)
                ):

                    log(
                        f"  {count:3d} | "
                        f"{shorten(value)}"
                    )

            # Save top 50 values for every column.
            for rank, (
                value,
                count
            ) in enumerate(
                counts.most_common(50),
                start=1
            ):

                value_rows.append(
                    {
                        "sheet": sheet_name,
                        "column": str(col),
                        "rank": rank,
                        "value": value,
                        "count": count,
                    }
                )

        log()
        log()

    # =========================================================
    # Read README
    # =========================================================

    log("=" * 80)
    log("README")
    log("=" * 80)

    if README.exists():

        text = README.read_text(
            encoding="utf-8",
            errors="replace"
        )

        log(text)

    else:

        log(
            "README file not found."
        )

    # =========================================================
    # Keyword scan for measurement-relevant columns
    # =========================================================

    keywords = [
        "metric",
        "measure",
        "productiv",
        "efficien",
        "performance",
        "satisfaction",
        "activity",
        "communication",
        "space",
        "time",
        "task",
        "quality",
        "review",
        "debug",
        "test",
        "integration",
        "delivery",
        "effort",
        "study",
        "method",
        "participant",
        "developer",
        "outcome",
    ]

    log()
    log("=" * 80)
    log("MEASUREMENT-RELEVANT COLUMN SCAN")
    log("=" * 80)

    matches = []

    for sheet_name, df in sheets.items():

        for col in df.columns:

            normalized = (
                str(col)
                .strip()
                .lower()
            )

            matched = [
                keyword
                for keyword in keywords
                if keyword in normalized
            ]

            if matched:

                matches.append(
                    (
                        sheet_name,
                        str(col),
                        ", ".join(matched),
                    )
                )

    if matches:

        for sheet, col, matched in matches:

            log(
                f"{sheet} | {col} "
                f"| keywords: {matched}"
            )

    else:

        log(
            "No measurement-related column "
            "names detected automatically."
        )

    # =========================================================
    # Save outputs
    # =========================================================

    structure_df = pd.DataFrame(
        structure_rows
    )

    completeness_df = pd.DataFrame(
        completeness_rows
    )

    values_df = pd.DataFrame(
        value_rows
    )

    structure_df.to_csv(
        STRUCTURE_FILE,
        index=False
    )

    completeness_df.to_csv(
        COMPLETENESS_FILE,
        index=False
    )

    values_df.to_csv(
        VALUES_FILE,
        index=False
    )

    LOG_FILE.write_text(
        "\n".join(log_lines),
        encoding="utf-8"
    )

    log()
    log("=" * 80)
    log("OUTPUT FILES")
    log("=" * 80)

    print(
        f"Audit log:          {LOG_FILE}"
    )

    print(
        f"Workbook structure: {STRUCTURE_FILE}"
    )

    print(
        f"Completeness:       {COMPLETENESS_FILE}"
    )

    print(
        f"Value inventory:    {VALUES_FILE}"
    )


if __name__ == "__main__":
    main()
